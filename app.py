"""WeChat Region Desk — local wizard for tagging PC WeChat contacts by region."""

from __future__ import annotations

import json
import os
import queue
import sys
import threading
import time
from datetime import datetime
from pathlib import Path
from typing import Any

import httpx
from fastapi import FastAPI, HTTPException
from fastapi.responses import FileResponse, Response, StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field

from axioxmedia import (
    AIO_BRAND,
    AIO_SOFTWARE_NAME_EN,
    AIO_SOFTWARE_NAME_ZH,
    aio_logo_png,
    aio_watermark,
    apply_hwnd_icon,
    axiox_window_title,
)
from utm_beacon import schedule_utm_beacon
from calibrate import detect_targets, image_to_screen, snap_point
from wechat_ops import (
    apply_tags,
    counts_from_rows,
    dump_tree,
    plan_from_rows,
    scan_contacts,
    snapshot_manager,
    test_click,
    window_status,
)

APP_VERSION = "1.9.3"
PRODUCT_EN = "WeChat Region Desk"
PRODUCT_ZH = "微信地区打标台"
SLUG = "wechat-region-desk"


def app_root() -> Path:
    if getattr(sys, "frozen", False) and hasattr(sys, "_MEIPASS"):
        return Path(sys._MEIPASS)
    return Path(__file__).resolve().parent


def runtime_dir() -> Path:
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent


ROOT = app_root()
STATIC = ROOT / "static"
RUN = runtime_dir()
OUTPUT = RUN / "output"
PREFS_PATH = RUN / f"{SLUG}-prefs.json"
LOG_FILE = RUN / "deploy_desk.log"
CONTACTS_PATH = OUTPUT / "contacts_region.json"
TREE_PATH = OUTPUT / "ui_tree.txt"

DEFAULT_PREFS = {
    "remember": True,
    "manager_window_name": "通讯录管理",
    "match_hongkong_keywords": ["中国香港", "香港", "Hong Kong", "HongKong", "HK"],
    "empty_region_tag": "未填写地区",
    "tag_prefix": "",
    "max_tag_name_length": 16,
    "scan_pause_seconds": 0.35,
    "click_pause_seconds": 0.25,
    "scan_limit": 15,
    "scan_all": False,
    "only_hk": False,
    "wechat_version": "4.1.13.65",
    "calibrate": {},
}

OUTPUT.mkdir(parents=True, exist_ok=True)

app = FastAPI(title=PRODUCT_ZH, version=APP_VERSION)
app.mount("/assets", StaticFiles(directory=STATIC), name="assets")

DESK_WINDOW = None
_job_lock = threading.Lock()
_job: dict[str, Any] = {
    "running": False,
    "kind": "",
    "q": queue.Queue(),
    "stop": threading.Event(),
    "result": None,
    "error": "",
}


class PrefsBody(BaseModel):
    remember: bool = True
    manager_window_name: str = "通讯录管理"
    match_hongkong_keywords: list[str] = Field(default_factory=list)
    empty_region_tag: str = "未填写地区"
    tag_prefix: str = ""
    max_tag_name_length: int = 16
    scan_pause_seconds: float = 0.35
    click_pause_seconds: float = 0.25
    scan_limit: int = 15
    scan_all: bool = False
    only_hk: bool = False
    wechat_version: str = "4.1.13.65"
    calibrate: dict[str, Any] = Field(default_factory=dict)


class ContactsBody(BaseModel):
    rows: list[dict[str, Any]]


def write_log(message: str) -> None:
    try:
        with LOG_FILE.open("a", encoding="utf-8") as fh:
            fh.write(f"{datetime.now().isoformat(timespec='seconds')} {message}\n")
    except Exception:
        pass


def load_prefs() -> dict:
    data = dict(DEFAULT_PREFS)
    if PREFS_PATH.exists():
        try:
            data.update(json.loads(PREFS_PATH.read_text(encoding="utf-8")))
        except Exception:
            pass
    return data


def save_prefs(data: dict) -> dict:
    merged = load_prefs()
    merged.update(data)
    if merged.get("remember", True):
        PREFS_PATH.write_text(json.dumps(merged, ensure_ascii=False, indent=2), encoding="utf-8")
    return merged


def load_rows() -> list[dict]:
    if not CONTACTS_PATH.exists():
        return []
    try:
        payload = json.loads(CONTACTS_PATH.read_text(encoding="utf-8"))
        return payload.get("rows") or []
    except Exception:
        return []


def save_rows(rows: list[dict]) -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    payload = {
        "saved_at": datetime.now().isoformat(timespec="seconds"),
        "rows": rows,
        "plan": plan_from_rows(rows),
        "counts": counts_from_rows(rows),
    }
    CONTACTS_PATH.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        wb = Workbook()
        ws = wb.active
        ws.title = "contacts"
        headers = ["nickname", "region_raw", "tag_name", "is_hongkong", "status", "note"]
        ws.append(headers)
        fill = PatternFill("solid", fgColor="1F4E3D")
        font = Font(color="FFFFFF", bold=True)
        for col in range(1, 7):
            ws.cell(1, col).fill = fill
            ws.cell(1, col).font = font
        hk = PatternFill("solid", fgColor="C6EFCE")
        for row in rows:
            ws.append([row.get(h, "") for h in headers])
            if row.get("is_hongkong"):
                for col in range(1, 7):
                    ws.cell(ws.max_row, col).fill = hk
        xlsx = OUTPUT / "contacts_region.xlsx"
        wb.save(xlsx)
        plan_wb = Workbook()
        pws = plan_wb.active
        pws.title = "tag_plan"
        pws.append(["tag_name", "count", "sample"])
        for item in payload["plan"]:
            pws.append([item["tag_name"], item["count"], item["sample"]])
        plan_wb.save(OUTPUT / "tag_plan.xlsx")
    except Exception as exc:  # noqa: BLE001
        write_log(f"xlsx skipped: {exc}")


@app.get("/brand/logo.png")
def brand_logo():
    return Response(content=aio_logo_png(), media_type="image/png")


@app.get("/favicon.ico")
def favicon():
    return Response(content=aio_logo_png(), media_type="image/png")


@app.get("/")
def index():
    return FileResponse(STATIC / "index.html")


@app.get("/api/defaults")
def api_defaults():
    prefs = load_prefs()
    rows = load_rows()
    return {
        "version": APP_VERSION,
        "product_zh": PRODUCT_ZH,
        "product_en": PRODUCT_EN,
        "brand": AIO_BRAND,
        "watermark": aio_watermark(),
        "prefs": prefs,
        "status": window_status(prefs.get("manager_window_name") or "通讯录管理"),
        "counts": counts_from_rows(rows),
        "plan": plan_from_rows(rows),
        "rows_preview": rows[:80],
        "rows_total": len(rows),
        "output_dir": str(OUTPUT),
        "axioxmedia": AIO_SOFTWARE_NAME_EN,
    }


@app.get("/api/prefs")
def api_prefs_get():
    return load_prefs()


@app.post("/api/prefs")
def api_prefs_post(body: PrefsBody):
    return save_prefs(body.model_dump())


@app.get("/api/status")
def api_status():
    prefs = load_prefs()
    st = window_status(prefs.get("manager_window_name") or "通讯录管理")
    st["job"] = {"running": _job["running"], "kind": _job["kind"], "error": _job["error"]}
    return st


@app.post("/api/inspect")
def api_inspect():
    prefs = load_prefs()
    st = window_status(prefs.get("manager_window_name") or "通讯录管理")
    if not st.get("found"):
        raise HTTPException(409, "通讯录管理窗口未打开")
    tree = dump_tree(prefs.get("manager_window_name") or "通讯录管理")
    OUTPUT.mkdir(parents=True, exist_ok=True)
    TREE_PATH.write_text(tree, encoding="utf-8")
    return {"ok": True, "status": st, "tree": tree[:12000], "path": str(TREE_PATH)}


class TestClickBody(BaseModel):
    x: int
    y: int


class SnapPointBody(BaseModel):
    mode: str
    img_x: int
    img_y: int


@app.post("/api/calibrate/snapshot")
def api_calibrate_snapshot():
    dest = OUTPUT / "calibrate.png"
    hidden = False
    try:
        if DESK_WINDOW is not None:
            DESK_WINDOW.hide()
            hidden = True
            time.sleep(0.4)
        info = snapshot_manager(dest)
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(409, str(exc)) from exc
    finally:
        if hidden and DESK_WINDOW is not None:
            try:
                DESK_WINDOW.show()
            except Exception:
                pass
    info["url"] = "/api/calibrate/image?ts=" + str(int(time.time()))
    (OUTPUT / "calibrate.json").write_text(json.dumps(info, ensure_ascii=False), encoding="utf-8")
    return info


@app.get("/api/calibrate/image")
def api_calibrate_image():
    path = OUTPUT / "calibrate.png"
    if not path.exists():
        raise HTTPException(404, "还没有截图")
    return FileResponse(path, media_type="image/png")


@app.post("/api/calibrate/test-click")
def api_calibrate_test(body: TestClickBody):
    try:
        test_click(body.x, body.y)
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(409, str(exc)) from exc
    return {"ok": True, "x": body.x, "y": body.y}


def _snap_meta() -> dict:
    path = OUTPUT / "calibrate.json"
    if not path.exists():
        raise HTTPException(404, "还没有截图")
    return json.loads(path.read_text(encoding="utf-8"))


@app.post("/api/calibrate/snap-point")
def api_snap_point(body: SnapPointBody):
    png = OUTPUT / "calibrate.png"
    if not png.exists():
        raise HTTPException(404, "还没有截图")
    meta = _snap_meta()
    try:
        hit = snap_point(png, body.mode, int(body.img_x), int(body.img_y))
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(409, str(exc)) from exc
    sx, sy = image_to_screen(meta, hit["img_x"], hit["img_y"])
    hit["screen_x"] = sx
    hit["screen_y"] = sy
    return hit


def _emit(kind: str, level: str, message: str, extra: dict | None = None) -> None:
    event = {"event": kind, "level": level, "message": message}
    if extra:
        event.update(extra)
    _job["q"].put(event)


def _run_scan(cfg: dict) -> None:
    rows_acc: list[dict] = []

    def log(msg: str) -> None:
        write_log(msg)
        _emit("log", "info", msg)

    try:
        rows = scan_contacts(
            cfg,
            log=log,
            should_stop=lambda: _job["stop"].is_set(),
            on_row=lambda rec: rows_acc.append(rec) or _emit("row", "info", rec.get("nickname") or "", {"row": rec}),
        )
        save_rows(rows)
        _job["result"] = {"counts": counts_from_rows(rows), "plan": plan_from_rows(rows)}
        _emit("done", "ok", "scan-complete", _job["result"])
    except Exception as exc:  # noqa: BLE001
        _job["error"] = str(exc)
        _emit("error", "error", str(exc))
    finally:
        _job["running"] = False


def _run_apply(cfg: dict) -> None:
    def log(msg: str) -> None:
        write_log(msg)
        _emit("log", "info", msg)

    try:
        rows = load_rows()
        if not rows:
            raise RuntimeError("还没有扫描结果")
        result = apply_tags(rows, cfg, log=log, should_stop=lambda: _job["stop"].is_set())
        _job["result"] = result
        _emit("done", "ok", "apply-complete", result)
    except Exception as exc:  # noqa: BLE001
        _job["error"] = str(exc)
        _emit("error", "error", str(exc))
    finally:
        _job["running"] = False


class ScanStartBody(BaseModel):
    scan_limit: int = 15
    scan_all: bool = False
    calibrate: dict[str, Any] = Field(default_factory=dict)


@app.post("/api/scan/start")
def api_scan_start(body: ScanStartBody | None = None):
    prefs = load_prefs()
    with _job_lock:
        if _job["running"]:
            raise HTTPException(409, "任务进行中")
        while not _job["q"].empty():
            _job["q"].get_nowait()
        _job.update({"running": True, "kind": "scan", "stop": threading.Event(), "result": None, "error": ""})
        cfg = dict(prefs)
        if body:
            cfg["scan_all"] = bool(body.scan_all)
            cfg["scan_limit"] = int(body.scan_limit)
            if body.calibrate:
                cfg["calibrate"] = body.calibrate
        cfg["limit"] = 0 if cfg.get("scan_all") else max(1, int(cfg.get("scan_limit") or 15))
        threading.Thread(target=_run_scan, args=(cfg,), daemon=True).start()
    return {"ok": True, "limit": cfg["limit"]}


@app.post("/api/apply/start")
def api_apply_start():
    prefs = load_prefs()
    with _job_lock:
        if _job["running"]:
            raise HTTPException(409, "任务进行中")
        while not _job["q"].empty():
            _job["q"].get_nowait()
        _job.update({"running": True, "kind": "apply", "stop": threading.Event(), "result": None, "error": ""})
        threading.Thread(target=_run_apply, args=(dict(prefs),), daemon=True).start()
    return {"ok": True}


@app.post("/api/job/stop")
def api_job_stop():
    _job["stop"].set()
    return {"ok": True}


@app.get("/api/scan/events")
@app.get("/api/apply/events")
@app.get("/api/job/events")
def api_job_events():
    def gen():
        while True:
            try:
                item = _job["q"].get(timeout=0.7)
            except queue.Empty:
                yield "event: ping\ndata: {}\n\n"
                if not _job["running"] and _job["q"].empty():
                    break
                continue
            yield f"event: {item.get('event', 'log')}\ndata: {json.dumps(item, ensure_ascii=False)}\n\n"
            if item.get("event") in {"done", "error"}:
                break

    return StreamingResponse(gen(), media_type="text/event-stream")


@app.get("/api/contacts")
def api_contacts():
    rows = load_rows()
    return {
        "rows": rows,
        "plan": plan_from_rows(rows),
        "counts": counts_from_rows(rows),
        "output_dir": str(OUTPUT),
    }


@app.post("/api/contacts/save")
def api_contacts_save(body: ContactsBody):
    save_rows(body.rows)
    return {"ok": True, "counts": counts_from_rows(body.rows), "plan": plan_from_rows(body.rows)}


@app.get("/api/export.xlsx")
def api_export_xlsx():
    path = OUTPUT / "contacts_region.xlsx"
    if not path.exists():
        raise HTTPException(404, "还没有扫描结果")
    return FileResponse(path, filename="contacts_region.xlsx")


def show_error(message: str) -> None:
    if os.name == "nt":
        try:
            import ctypes

            ctypes.windll.user32.MessageBoxW(0, message, PRODUCT_EN, 0x10)
            return
        except Exception:
            pass
    print(message, file=sys.stderr)


def _free_port(preferred: int = 8787) -> int:
    import socket

    for port in (preferred, 8788, 8789, 8790, 0):
        sock = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
        try:
            sock.setsockopt(socket.SOL_SOCKET, socket.SO_REUSEADDR, 1)
            sock.bind(("127.0.0.1", port))
            chosen = int(sock.getsockname()[1])
        except OSError:
            chosen = -1
        finally:
            sock.close()
        if chosen > 0:
            return chosen
    raise RuntimeError("没有可用的本地端口")


def ensure_stdio() -> None:
    if sys.stdout is None:
        sys.stdout = LOG_FILE.open("a", encoding="utf-8")
    if sys.stderr is None:
        sys.stderr = LOG_FILE.open("a", encoding="utf-8")


def run_server(host: str, port: int, reload: bool = False) -> None:
    import uvicorn

    ensure_stdio()
    if reload:
        uvicorn.run(app, host=host, port=port, reload=True, log_level="warning", log_config=None)
        return
    config = uvicorn.Config(
        app,
        host=host,
        port=port,
        log_level="warning",
        log_config=None,
        lifespan="on",
        access_log=False,
    )
    server = uvicorn.Server(config)
    server.install_signal_handlers = False
    server.run()


def wait_ready(url: str, server_error: list[str], timeout: float = 30.0) -> None:
    deadline = time.time() + timeout
    while time.time() < deadline:
        if server_error:
            raise RuntimeError(server_error[0])
        try:
            with httpx.Client(timeout=0.8, trust_env=False) as http:
                if http.get(url).status_code < 500:
                    return
        except httpx.HTTPError:
            time.sleep(0.2)
    extra = f"\n服务线程错误：{server_error[0]}" if server_error else ""
    raise RuntimeError(f"本地服务启动超时：{url}{extra}\n日志：{LOG_FILE}")


def run_desktop() -> None:
    import traceback
    import webbrowser

    try:
        from wechat_ops import _dpi_aware

        _dpi_aware()
    except Exception:
        pass
    write_log(f"start frozen={getattr(sys, 'frozen', False)} meipass={getattr(sys, '_MEIPASS', '')}")
    port = _free_port()
    url = f"http://127.0.0.1:{port}"
    write_log(f"bind {url}")
    server_error: list[str] = []

    def _serve() -> None:
        try:
            run_server("127.0.0.1", port, reload=False)
        except Exception:
            server_error.append(traceback.format_exc())
            write_log(server_error[-1])

    thread = threading.Thread(target=_serve, name="uvicorn", daemon=True)
    thread.start()
    wait_ready(f"{url}/api/defaults", server_error)
    schedule_utm_beacon(product_en=PRODUCT_EN, version=APP_VERSION, log=write_log)
    try:
        import webview

        global DESK_WINDOW
        window = webview.create_window(
            title=axiox_window_title(PRODUCT_ZH, PRODUCT_EN),
            url=url,
            width=1280,
            height=860,
            min_size=(960, 680),
            background_color="#0b0d12",
        )
        DESK_WINDOW = window

        def paint_chrome(_=None) -> None:
            if os.name != "nt":
                return
            try:
                import ctypes

                hwnd = int(window.native.Handle.ToInt32())
                apply_hwnd_icon(hwnd)
                value = ctypes.c_int(1)
                for attr in (20, 19):
                    ctypes.windll.dwmapi.DwmSetWindowAttribute(
                        hwnd, attr, ctypes.byref(value), ctypes.sizeof(value)
                    )
            except Exception as exc:
                write_log(f"dark titlebar skipped: {exc}")

        try:
            window.events.shown += paint_chrome
        except Exception:
            pass
        webview.start()
        return
    except Exception:
        write_log(traceback.format_exc())
        webbrowser.open(url)
        while thread.is_alive():
            thread.join(timeout=0.5)


if __name__ == "__main__":
    import multiprocessing
    import traceback

    multiprocessing.freeze_support()
    ensure_stdio()
    try:
        desktop = "--web" not in sys.argv and os.environ.get("DEPLOY_DESK_WEB") != "1"
        if desktop:
            run_desktop()
        else:
            run_server("127.0.0.1", _free_port(8787), reload=not getattr(sys, "frozen", False))
    except Exception:
        show_error("启动失败：\n\n" + traceback.format_exc() + f"\n\n日志文件：{LOG_FILE}")
        raise
